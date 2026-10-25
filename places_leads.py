#!/usr/bin/env python3
"""
places_leads.py - Build a scored list of local business leads from the Google Places API (New).

For each category (restaurants, clinics, clothing stores, ...) it searches your area,
collects name, phone, website, rating, reviews etc., scores each business for
"website / app opportunity", and saves everything to an Excel file sorted by score.

Usage:
    cp .env.example .env      # then put your API key (and default location) in .env
    python places_leads.py

    # Or override any setting on the command line:
    python places_leads.py --location "Your City" --lat 26.76 --lng 83.37 --radius 8000

    # Search a grid of smaller circles to find many more businesses (costs more API calls):
    python places_leads.py --location "Your City" --lat 26.76 --lng 83.37 --radius 8000 --grid 3

    # Only some categories:
    python places_leads.py --location "Your City" --categories "restaurants,clinics,gyms"
"""

import argparse
import csv
import math
import os
import sys
import time

import requests

API_URL = "https://places.googleapis.com/v1/places:searchText"

# Fields we ask Google for. Phone + website + rating put the request in a higher
# pricing tier; check current pricing at https://developers.google.com/maps/billing-and-pricing
FIELD_MASK = ",".join([
    "places.id",
    "places.displayName",
    "places.formattedAddress",
    "places.nationalPhoneNumber",
    "places.internationalPhoneNumber",
    "places.websiteUri",
    "places.rating",
    "places.userRatingCount",
    "places.primaryTypeDisplayName",
    "places.googleMapsUri",
    "places.businessStatus",
    "places.location",
    "nextPageToken",
])

DEFAULT_CATEGORIES = [
    "restaurants", "cafes", "sweet shops", "hospitals", "clinics", "dentists",
    "pharmacies", "clothing stores", "boutiques", "salons", "gyms",
    "coaching centers", "schools", "hotels", "jewellers", "electronics shops",
    "furniture shops", "real estate agents",
]

# Categories that gain most from booking / menu / ordering features (+10 points)
HIGH_VALUE_WORDS = [
    "restaurant", "cafe", "café", "hospital", "clinic", "dentist", "doctor",
    "coaching", "school", "salon", "spa", "gym", "fitness", "hotel",
]

SOCIAL_DOMAINS = ["instagram.com", "facebook.com", "fb.com", "wa.me", "whatsapp.com"]


# --------------------------------------------------------------------------- API

def search_text(api_key, query, lat=None, lng=None, radius=None, max_pages=3):
    """Run one Text Search query and follow pagination (max 20 results/page, 60 total)."""
    headers = {
        "Content-Type": "application/json",
        "X-Goog-Api-Key": api_key,
        "X-Goog-FieldMask": FIELD_MASK,
    }
    body = {"textQuery": query, "pageSize": 20}
    if lat is not None and lng is not None:
        body["locationBias"] = {
            "circle": {
                "center": {"latitude": lat, "longitude": lng},
                "radius": float(min(radius or 5000, 50000)),
            }
        }

    results = []
    for _ in range(max_pages):
        for attempt in range(3):
            resp = requests.post(API_URL, headers=headers, json=body, timeout=30)
            if resp.status_code == 429 or resp.status_code >= 500:
                time.sleep(2 ** attempt)
                continue
            break
        if resp.status_code != 200:
            print(f"  ! Error {resp.status_code} for '{query}': {resp.text[:300]}", file=sys.stderr)
            break

        data = resp.json()
        results.extend(data.get("places", []))
        token = data.get("nextPageToken")
        if not token:
            break
        body = {**body, "pageToken": token}
        time.sleep(1)  # small pause before requesting the next page
    return results


def grid_points(lat, lng, radius_m, n):
    """Split a circle into an n x n grid of smaller circles (to get past the 60-result cap)."""
    if n <= 1:
        return [(lat, lng, radius_m)]
    step = (2 * radius_m) / n
    cell_radius = step * 0.75  # overlap a bit so nothing falls between cells
    lat_per_m = 1 / 111_320
    lng_per_m = 1 / (111_320 * math.cos(math.radians(lat)))
    points = []
    for i in range(n):
        for j in range(n):
            dy = -radius_m + step * (i + 0.5)
            dx = -radius_m + step * (j + 0.5)
            points.append((lat + dy * lat_per_m, lng + dx * lng_per_m, cell_radius))
    return points


# ----------------------------------------------------------------------- scoring

def score_business(row):
    """Opportunity score, using the rules from the action-plan doc."""
    score = 0
    reasons = []
    website = (row["website"] or "").lower()
    reviews = row["reviews"] or 0
    rating = row["rating"] or 0
    category = (row["category"] + " " + row["searched_for"]).lower()
    social_only = any(d in website for d in SOCIAL_DOMAINS)

    if not website or social_only:
        score += 40
        reasons.append("no website")
        if reviews >= 100:
            score += 30
            reasons.append("100+ reviews, no website")
        if social_only:
            score += 15
            reasons.append("social media only")
    else:
        # Has a real website: could still be a redesign lead (check it by hand)
        reasons.append("has website - check if old/broken (+15 if so)")

    if rating >= 4.0:
        score += 10
        reasons.append("rating 4+")
    if any(w in category for w in HIGH_VALUE_WORDS):
        score += 10
        reasons.append("high-value category")
    if reviews < 10:
        score -= 10
        reasons.append("under 10 reviews")

    if score >= 70:
        group = "Hot"
    elif score >= 40:
        group = "Warm"
    else:
        group = "Later"
    return score, group, "; ".join(reasons)


def to_row(place, searched_for):
    return {
        "place_id": place.get("id", ""),
        "name": (place.get("displayName") or {}).get("text", ""),
        "category": (place.get("primaryTypeDisplayName") or {}).get("text", ""),
        "searched_for": searched_for,
        "address": place.get("formattedAddress", ""),
        "phone": place.get("nationalPhoneNumber") or place.get("internationalPhoneNumber", ""),
        "website": place.get("websiteUri", ""),
        "rating": place.get("rating"),
        "reviews": place.get("userRatingCount", 0),
        "status": place.get("businessStatus", ""),
        "maps_link": place.get("googleMapsUri", ""),
        "lat": (place.get("location") or {}).get("latitude"),
        "lng": (place.get("location") or {}).get("longitude"),
    }


# ------------------------------------------------------------------------ output

COLUMNS = [
    "score", "group", "name", "category", "searched_for", "phone", "website",
    "rating", "reviews", "address", "maps_link", "score_reasons", "status",
    "lead_status", "last_contact", "next_follow_up", "notes", "deal_value_inr",
    "place_id", "lat", "lng",
]


def save_csv(rows, path):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNS, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def save_xlsx(rows, path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError:
        print("openpyxl not installed - skipping Excel file (CSV still saved).")
        return False

    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    ws.append(COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
    fills = {"Hot": "F8CBAD", "Warm": "FFE699", "Later": "E2EFDA"}
    for r in rows:
        ws.append([r.get(c, "") for c in COLUMNS])
        fill = PatternFill("solid", fgColor=fills[r["group"]])
        ws.cell(row=ws.max_row, column=2).fill = fill

    # Dropdown for tracking each lead
    dv = DataValidation(
        type="list",
        formula1='"New,Contacted,Demo shown,Interested,Won,Lost,Later"',
        allow_blank=True,
    )
    ws.add_data_validation(dv)
    col = COLUMNS.index("lead_status") + 1
    letter = ws.cell(row=1, column=col).column_letter
    dv.add(f"{letter}2:{letter}{max(ws.max_row, 2)}")

    widths = {"name": 32, "address": 45, "website": 30, "maps_link": 30,
              "score_reasons": 45, "notes": 30, "category": 20, "searched_for": 18}
    for i, c in enumerate(COLUMNS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = widths.get(c, 13)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    # Summary sheet
    s = wb.create_sheet("Summary")
    s.append(["Category", "Total", "No website", "Hot", "Warm"])
    for c in s[1]:
        c.font = Font(bold=True)
    cats = sorted({r["searched_for"] for r in rows})
    for cat in cats:
        sub = [r for r in rows if r["searched_for"] == cat]
        s.append([
            cat, len(sub),
            sum(1 for r in sub if "no website" in r["score_reasons"]),
            sum(1 for r in sub if r["group"] == "Hot"),
            sum(1 for r in sub if r["group"] == "Warm"),
        ])
    s.column_dimensions["A"].width = 22
    wb.save(path)
    return True


# -------------------------------------------------------------------------- main

def load_env(path=".env"):
    """Load KEY=VALUE lines from a .env file into os.environ (existing env vars win)."""
    if not os.path.exists(path):
        return
    with open(path, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, value = line.split("=", 1)
            value = value.strip().strip('"').strip("'")
            if value:
                os.environ.setdefault(key.strip(), value)


def env_float(name):
    v = os.environ.get(name)
    return float(v) if v else None


def main():
    load_env(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".env"))
    load_env(".env")

    p = argparse.ArgumentParser(description="Get scored local business leads from Google Places API.")
    p.add_argument("--location", default=os.environ.get("LOCATION"),
                   help='Area name added to each search, e.g. "Civil Lines, Your City" (or LOCATION in .env)')
    p.add_argument("--lat", type=float, default=env_float("LAT"), help="Centre latitude (or LAT in .env)")
    p.add_argument("--lng", type=float, default=env_float("LNG"), help="Centre longitude (or LNG in .env)")
    p.add_argument("--radius", type=float, default=env_float("RADIUS") or 5000,
                   help="Search radius in metres (default 5000, max 50000)")
    p.add_argument("--grid", type=int, default=int(os.environ.get("GRID") or 1),
                   help="Split area into N x N cells for more results (default 1)")
    p.add_argument("--categories", default=os.environ.get("CATEGORIES"),
                   help="Comma-separated categories (default: built-in list of 18)")
    p.add_argument("--output", default="business_leads", help="Output file name without extension")
    p.add_argument("--include-closed", action="store_true", help="Keep permanently/temporarily closed businesses")
    p.add_argument("--api-key", default=os.environ.get("GOOGLE_MAPS_API_KEY"), help="Or set GOOGLE_MAPS_API_KEY")
    args = p.parse_args()

    if not args.api_key:
        sys.exit("Missing API key. Add GOOGLE_MAPS_API_KEY to .env (see .env.example) or pass --api-key.")
    if not args.location:
        sys.exit("Missing location. Add LOCATION to .env or pass --location.")
    if args.grid > 1 and (args.lat is None or args.lng is None):
        sys.exit("--grid needs --lat and --lng.")

    categories = [c.strip() for c in args.categories.split(",")] if args.categories else DEFAULT_CATEGORIES
    if args.lat is not None and args.lng is not None:
        points = grid_points(args.lat, args.lng, args.radius, args.grid)
    else:
        points = [(None, None, None)]

    print(f"Searching {len(categories)} categories x {len(points)} area(s) in {args.location} ...")
    seen = {}
    for cat in categories:
        before = len(seen)
        for lat, lng, rad in points:
            for place in search_text(args.api_key, f"{cat} in {args.location}", lat, lng, rad):
                pid = place.get("id")
                if pid and pid not in seen:
                    seen[pid] = to_row(place, cat)
        print(f"  {cat:<22} +{len(seen) - before} new")

    rows = list(seen.values())
    if not args.include_closed:
        rows = [r for r in rows if r["status"] in ("", "OPERATIONAL")]

    for r in rows:
        r["score"], r["group"], r["score_reasons"] = score_business(r)
        r["lead_status"] = "New"
    rows.sort(key=lambda r: (-r["score"], -(r["reviews"] or 0)))

    save_csv(rows, args.output + ".csv")
    made_xlsx = save_xlsx(rows, args.output + ".xlsx")

    hot = sum(1 for r in rows if r["group"] == "Hot")
    warm = sum(1 for r in rows if r["group"] == "Warm")
    no_site = sum(1 for r in rows if "no website" in r["score_reasons"])
    print(f"\nDone: {len(rows)} businesses | {no_site} without website | {hot} hot | {warm} warm")
    print(f"Saved {args.output}.csv" + (f" and {args.output}.xlsx" if made_xlsx else ""))


if __name__ == "__main__":
    main()
